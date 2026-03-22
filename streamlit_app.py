import math
import re
from datetime import datetime, timedelta, timezone
from io import BytesIO

import pandas as pd
import streamlit as st
from openpyxl import load_workbook

st.set_page_config(page_title="Avito Master Tool Final + Preview", layout="wide")
st.title("Avito Master Tool Final + Preview")

col1, col2, col3, col4 = st.columns(4)
with col1:
    margin_percent = st.number_input("Изменение цены (%)", value=-12.0, step=1.0, format="%.1f")
with col2:
    active_days = st.number_input("DateEnd для активных (дней вперёд)", value=30, step=1, min_value=1)
with col3:
    close_days_back = st.number_input("DateEnd для закрытых (дней назад)", value=7, step=1, min_value=1)
with col4:
    utc_offset_hours = st.number_input("UTC offset", value=3, step=1)

price_file = st.file_uploader("1) Прайс", type=["xlsx"])
avito_file = st.file_uploader("2) Файл Авито", type=["xlsx"])
stock_file = st.file_uploader("3) Файл остатков", type=["xlsx"])

ARTICLE_RE = re.compile(r"\b[A-Z0-9][A-Z0-9\-]{3,}\b")

NON_ORIGINAL_MARKERS = [
    "compatible",
    "совместим",
    "совместимый",
    "аналог",
    "noname",
    "no name",

    "g&g",
    "gg",
    "cet",
    "cet group",
    "aquamarine",
    "cactus",
    "sakura",
    "dataproducts",
    "retech",
    "uniton",
    "uniton premium",
    "hi-black",
    "hiblack",
    "profiline",
    "colortek",
    "7q",
    "nv print",
    "nv-print",
    "tonex",
    "mse",
    "freecolor",
    "kodak",
    "integral",
    "static control",
    "solnce",
    "hyb toner",
    "t2",
    "easy print",
]

NON_ORIGINAL_AVITO_MARKERS = NON_ORIGINAL_MARKERS.copy()


def clean(x):
    if x is None:
        return ""
    if isinstance(x, float) and pd.isna(x):
        return ""
    return str(x).strip()


def norm_article(x):
    s = clean(x).upper()
    s = s.replace(".JPG", "").replace(".JPEG", "").replace(".PNG", "").replace(".WEBP", "")
    s = re.sub(r"\s+", "", s)
    return s


def round_up_100(value):
    return int(math.ceil(float(value) / 100.0) * 100)


def current_dt(offset_hours: int):
    tz = timezone(timedelta(hours=int(offset_hours)))
    return datetime.now(tz).replace(microsecond=0)


def fmt_dt(dt: datetime):
    return dt.isoformat()


def parse_numeric(value):
    if value is None or (isinstance(value, float) and pd.isna(value)):
        return None
    if isinstance(value, (int, float)):
        return float(value)
    s = str(value).strip().replace(" ", "").replace("\xa0", "").replace(",", ".")
    if not s:
        return None
    try:
        return float(s)
    except Exception:
        return None


def candidate_tokens(text):
    txt = clean(text).upper()
    txt = re.sub(r"(\d{3})\s*/\s*R\s*([0-9]{4,5})", r"\1R\2", txt)
    txt = re.sub(r"(\d{3})\s*/\s*R0*([0-9]{4,5})", r"\1R\2", txt)
    txt = re.sub(r"№\s*([A-Z0-9\-]+)", r" \1 ", txt)
    txt = re.sub(r"<[^>]+>", " ", txt)

    found = ARTICLE_RE.findall(txt)
    extra = []
    for tok in found:
        if "/" in tok:
            extra.extend(tok.split("/"))
    found.extend(extra)

    uniq = []
    for tok in found:
        tok = tok.strip(".,;:()[]{}<>\"'")
        tok = norm_article(tok)
        if tok and tok not in uniq:
            uniq.append(tok)

    return sorted(uniq, key=lambda s: (-len(s), s))


def contains_non_original_marker(text: str, markers) -> bool:
    s = clean(text).lower().replace("ё", "е")
    return any(marker in s for marker in markers)


def is_non_original_price_row(row_text: str, manufacturer: str) -> bool:
    txt = clean(row_text).lower().replace("ё", "е")
    mfr = clean(manufacturer).lower().replace("ё", "е")

    if contains_non_original_marker(txt, NON_ORIGINAL_MARKERS):
        return True
    if contains_non_original_marker(mfr, NON_ORIGINAL_MARKERS):
        return True

    return False


def is_original_avito_row(title: str, desc: str) -> bool:
    full = f"{clean(title)} {clean(desc)}".lower().replace("ё", "е")
    return not contains_non_original_marker(full, NON_ORIGINAL_AVITO_MARKERS)


def choose_article(title, desc, price_map):
    for source in [desc, title, f"{title} {desc}"]:
        for tok in candidate_tokens(source):
            if tok in price_map:
                return tok
    return None


def find_col_index(headers, keywords):
    for i, h in enumerate(headers, start=1):
        hl = clean(h).lower()
        for kw in keywords:
            if kw in hl:
                return i
    return None


def load_price_map(file_obj):
    df = pd.read_excel(file_obj)
    df.columns = [clean(c) for c in df.columns]

    article_col = None
    price_col = None
    stock_col = None
    manufacturer_col = None
    nomenclature_col = None

    for c in df.columns:
        cl = c.lower()
        if article_col is None and "артик" in cl:
            article_col = c
        if price_col is None and cl == "цена":
            price_col = c
        if stock_col is None and "свобод" in cl:
            stock_col = c
        if manufacturer_col is None and "производитель" in cl:
            manufacturer_col = c
        if nomenclature_col is None and cl == "номенклатура":
            nomenclature_col = c

    if price_col is None and len(df.columns) >= 6:
        price_col = df.columns[5]

    if article_col is None or price_col is None:
        raise ValueError("В прайсе не найдены обязательные колонки: Артикул и Цена (или колонка F).")

    price_map = {}
    stock_map = {}
    skipped_non_original = 0
    skipped_empty_article = 0

    for _, row in df.iterrows():
        article_raw = clean(row[article_col])
        art = norm_article(article_raw)
        if not art:
            skipped_empty_article += 1
            continue

        manufacturer = clean(row[manufacturer_col]) if manufacturer_col else ""
        nomenclature = clean(row[nomenclature_col]) if nomenclature_col else ""
        full_row_text = " ".join(clean(v) for v in row.tolist())

        if is_non_original_price_row(
            row_text=f"{full_row_text} {nomenclature}",
            manufacturer=manufacturer
        ):
            skipped_non_original += 1
            continue

        price = parse_numeric(row[price_col])
        if price is not None:
            price_map[art] = price

        stock = 0
        if stock_col is not None:
            sv = parse_numeric(row[stock_col])
            stock = int(sv) if sv is not None else 0

        stock_map[art] = stock

    return price_map, stock_map, {
        "article_col": article_col,
        "price_col": price_col,
        "stock_col": stock_col,
        "manufacturer_col": manufacturer_col,
        "nomenclature_col": nomenclature_col,
        "rows": len(df),
        "original_positions_loaded": len(price_map),
        "skipped_non_original_rows": skipped_non_original,
        "skipped_empty_article_rows": skipped_empty_article,
    }


def calc_dateend_from_qty(qty, offset_hours, active_days, close_days_back):
    now_dt = current_dt(offset_hours)
    if qty <= 0:
        return fmt_dt(now_dt - timedelta(days=int(close_days_back)))
    return fmt_dt(now_dt + timedelta(days=int(active_days)))


def calc_dateend_from_status(status_text, offset_hours, active_days, close_days_back):
    s = clean(status_text).lower()
    now_dt = current_dt(offset_hours)

    if "снято" in s or "неактив" in s or "архив" in s:
        return fmt_dt(now_dt - timedelta(days=int(close_days_back)))
    return fmt_dt(now_dt + timedelta(days=int(active_days)))


def ensure_dateend_column(ws):
    headers = [clean(c.value) for c in ws[2]]

    dateend_col = None
    avito_dateend_col = None

    for i, h in enumerate(headers, start=1):
        hl = h.lower()
        if hl == "dateend":
            dateend_col = i
        if hl == "avitodateend":
            avito_dateend_col = i

    if dateend_col:
        return dateend_col

    if avito_dateend_col:
        ws.cell(row=2, column=avito_dateend_col).value = "DateEnd"
        return avito_dateend_col

    new_col = ws.max_column + 1
    ws.cell(row=2, column=new_col).value = "DateEnd"
    return new_col


def detect_sheet_columns(ws):
    headers = [clean(c.value) for c in ws[2]]

    cols = {
        "title_col": find_col_index(headers, ["название"]),
        "desc_col": find_col_index(headers, ["описание"]),
        "price_col": find_col_index(headers, ["цена"]),
        "status_col": find_col_index(headers, ["avitostatus"]),
        "ad_number_col": None,
    }

    for i, h in enumerate(headers, start=1):
        hl = clean(h).lower()
        if "номер объявления на авито" in hl:
            cols["ad_number_col"] = i
            break

    if cols["ad_number_col"] is None:
        cols["ad_number_col"] = find_col_index(headers, ["номер объявления"])

    cols["dateend_col"] = ensure_dateend_column(ws)
    return cols


def preview_changes(avito_file_obj, price_map, stock_map, margin_percent, offset_hours, active_days, close_days_back):
    wb = load_workbook(avito_file_obj, data_only=True)

    update_rows = []
    skipped_non_original = []
    new_stock_ids = []

    for ws in wb.worksheets:
        cols = detect_sheet_columns(ws)
        title_col = cols["title_col"]
        desc_col = cols["desc_col"]
        price_col = cols["price_col"]
        status_col = cols["status_col"]
        ad_number_col = cols["ad_number_col"]

        if not title_col:
            continue

        for row in range(5, ws.max_row + 1):
            title = ws.cell(row=row, column=title_col).value if title_col else ""
            desc = ws.cell(row=row, column=desc_col).value if desc_col else ""
            old_price = ws.cell(row=row, column=price_col).value if price_col else ""
            status_text = ws.cell(row=row, column=status_col).value if status_col else ""
            avito_id = clean(ws.cell(row=row, column=ad_number_col).value) if ad_number_col else ""

            if clean(title) == "":
                continue

            if not is_original_avito_row(title, desc):
                skipped_non_original.append({
                    "sheet": ws.title,
                    "row": row,
                    "avito_id": avito_id,
                    "title": clean(title),
                })
                continue

            article = choose_article(title, desc, price_map)

            if article:
                qty = int(stock_map.get(article, 0))
                new_price = round_up_100(price_map[article] * (1 + margin_percent / 100.0)) if article in price_map else old_price
                new_dateend = calc_dateend_from_qty(qty, offset_hours, active_days, close_days_back)
                action = "Закрыть" if qty <= 0 else "Оставить активным"
            else:
                qty = None
                new_price = old_price
                new_dateend = calc_dateend_from_status(status_text, offset_hours, active_days, close_days_back)
                action = "DateEnd по статусу"

            update_rows.append({
                "sheet": ws.title,
                "row": row,
                "avito_id": avito_id,
                "article": article or "",
                "title": clean(title),
                "old_price": old_price,
                "new_price": new_price,
                "qty": qty,
                "new_dateend": new_dateend,
                "action": action,
            })

    return pd.DataFrame(update_rows), pd.DataFrame(skipped_non_original)


def build_avito_article_map(avito_file_obj, price_map):
    wb = load_workbook(avito_file_obj, data_only=True)
    ad_to_article = {}
    matched = 0
    unmatched = 0
    skipped_non_original_avito = 0

    for ws in wb.worksheets:
        cols = detect_sheet_columns(ws)
        title_col = cols["title_col"]
        desc_col = cols["desc_col"]
        ad_number_col = cols["ad_number_col"]

        if not ad_number_col:
            continue

        for row in range(5, ws.max_row + 1):
            ad_number = ws.cell(row=row, column=ad_number_col).value
            if not ad_number:
                continue

            title = ws.cell(row=row, column=title_col).value if title_col else ""
            desc = ws.cell(row=row, column=desc_col).value if desc_col else ""

            if not is_original_avito_row(title, desc):
                skipped_non_original_avito += 1
                continue

            article = choose_article(title, desc, price_map)

            if article:
                ad_to_article[clean(ad_number)] = article
                matched += 1
            else:
                unmatched += 1

    return ad_to_article, {
        "matched": matched,
        "unmatched": unmatched,
        "skipped_non_original_avito": skipped_non_original_avito,
    }


def preview_new_stock_ids(stock_file_obj, avito_file_obj, price_map):
    ad_to_article, _ = build_avito_article_map(avito_file_obj, price_map)

    wb = load_workbook(stock_file_obj, data_only=True)
    ws = wb.active

    headers = [clean(c.value) for c in ws[1]]
    avito_id_col = find_col_index(headers, ["avitoid"])

    existing_ids = set()
    if avito_id_col:
        for row in range(2, ws.max_row + 1):
            avito_id = clean(ws.cell(row=row, column=avito_id_col).value)
            if avito_id:
                existing_ids.add(avito_id)

    new_ids = []
    for avito_id, article in ad_to_article.items():
        if avito_id not in existing_ids:
            new_ids.append({
                "avito_id": avito_id,
                "article": article,
            })

    return pd.DataFrame(new_ids)


def update_prices_and_dateend(avito_file_obj, price_map, stock_map, margin_percent, offset_hours, active_days, close_days_back):
    wb = load_workbook(avito_file_obj)
    updated_rows = 0
    updated_prices = 0
    dateend_set_from_qty = 0
    dateend_set_from_status = 0
    unmatched = 0
    skipped_non_original_avito = 0
    examples = []

    for ws in wb.worksheets:
        cols = detect_sheet_columns(ws)
        title_col = cols["title_col"]
        desc_col = cols["desc_col"]
        price_col = cols["price_col"]
        status_col = cols["status_col"]
        dateend_col = cols["dateend_col"]

        if not title_col:
            continue

        for row in range(5, ws.max_row + 1):
            title = ws.cell(row=row, column=title_col).value if title_col else ""
            desc = ws.cell(row=row, column=desc_col).value if desc_col else ""
            status_text = ws.cell(row=row, column=status_col).value if status_col else ""

            if clean(title) == "":
                continue

            if not is_original_avito_row(title, desc):
                skipped_non_original_avito += 1
                continue

            article = choose_article(title, desc, price_map)

            if article:
                qty = int(stock_map.get(article, 0))

                if article in price_map and price_col:
                    base_price = price_map[article]
                    new_price = round_up_100(base_price * (1 + margin_percent / 100.0))
                    ws.cell(row=row, column=price_col).value = new_price
                    updated_prices += 1

                new_dateend = calc_dateend_from_qty(
                    qty=qty,
                    offset_hours=offset_hours,
                    active_days=active_days,
                    close_days_back=close_days_back,
                )
                ws.cell(row=row, column=dateend_col).value = new_dateend
                dateend_set_from_qty += 1

                if status_col:
                    ws.cell(row=row, column=status_col).value = "Активно" if qty > 0 else "Снято с публикации"

                updated_rows += 1

                if len(examples) < 15:
                    examples.append({
                        "sheet": ws.title,
                        "row": row,
                        "article": article,
                        "mode": "from_qty",
                        "qty": qty,
                        "dateend": new_dateend,
                        "price": ws.cell(row=row, column=price_col).value if price_col else "",
                    })

            else:
                unmatched += 1

                fallback_dateend = calc_dateend_from_status(
                    status_text=status_text,
                    offset_hours=offset_hours,
                    active_days=active_days,
                    close_days_back=close_days_back,
                )
                ws.cell(row=row, column=dateend_col).value = fallback_dateend
                dateend_set_from_status += 1
                updated_rows += 1

                if len(examples) < 15:
                    examples.append({
                        "sheet": ws.title,
                        "row": row,
                        "article": "",
                        "mode": "from_status",
                        "status": clean(status_text),
                        "dateend": fallback_dateend,
                        "price": ws.cell(row=row, column=price_col).value if price_col else "",
                    })

    output = BytesIO()
    wb.save(output)
    output.seek(0)

    return output, {
        "updated_rows": updated_rows,
        "updated_prices": updated_prices,
        "dateend_set_from_qty": dateend_set_from_qty,
        "dateend_set_from_status": dateend_set_from_status,
        "unmatched": unmatched,
        "skipped_non_original_avito": skipped_non_original_avito,
        "examples": examples,
    }


def update_stock_only(stock_file_obj, avito_file_obj, price_map, stock_map):
    ad_to_article, ad_map_stats = build_avito_article_map(avito_file_obj, price_map)

    wb = load_workbook(stock_file_obj)
    ws = wb.active

    headers = [clean(c.value) for c in ws[1]]
    avito_id_col = find_col_index(headers, ["avitoid"])
    stock_col = find_col_index(headers, ["stock"])

    if not avito_id_col or not stock_col:
        raise ValueError("В файле остатков не найдены колонки AvitoId / Stock.")

    existing_ids = {}
    for row in range(2, ws.max_row + 1):
        avito_id = clean(ws.cell(row=row, column=avito_id_col).value)
        if avito_id:
            existing_ids[avito_id] = row

    updated_existing = 0
    added_new = 0
    zero_count = 0
    one_count = 0
    many_count = 0
    unmatched_article_count = 0

    def stock_payload(qty):
        if qty <= 0:
            return 0
        elif qty == 1:
            return 1
        else:
            return "*"

    for avito_id, row in existing_ids.items():
        article = ad_to_article.get(avito_id)
        if article:
            qty = int(stock_map.get(article, 0))
        else:
            qty = 0
            unmatched_article_count += 1

        stock_value = stock_payload(qty)
        ws.cell(row=row, column=stock_col).value = stock_value
        updated_existing += 1

        if stock_value == 0:
            zero_count += 1
        elif stock_value == 1:
            one_count += 1
        else:
            many_count += 1

    next_row = ws.max_row + 1
    for avito_id, article in ad_to_article.items():
        if avito_id in existing_ids:
            continue

        qty = int(stock_map.get(article, 0))
        stock_value = stock_payload(qty)

        ws.cell(row=next_row, column=avito_id_col).value = avito_id
        ws.cell(row=next_row, column=stock_col).value = stock_value
        next_row += 1
        added_new += 1

        if stock_value == 0:
            zero_count += 1
        elif stock_value == 1:
            one_count += 1
        else:
            many_count += 1

    output = BytesIO()
    wb.save(output)
    output.seek(0)

    return output, {
        "updated_existing": updated_existing,
        "added_new": added_new,
        "zero_count": zero_count,
        "one_count": one_count,
        "many_count": many_count,
        "unmatched_article_count": unmatched_article_count,
        "ad_map_stats": ad_map_stats,
    }


col_a, col_b, col_c, col_d = st.columns(4)
run_preview = col_a.button("Предпросмотр", use_container_width=True)
run_prices = col_b.button("Обновить цены + DateEnd", use_container_width=True)
run_stock = col_c.button("Обновить только Stock", use_container_width=True)
run_all = col_d.button("Обновить всё", use_container_width=True)

if run_preview or run_prices or run_stock or run_all:
    if not price_file or not avito_file:
        st.error("Для работы нужны минимум: прайс и файл Авито.")
        st.stop()

    try:
        price_map, stock_map, price_diag = load_price_map(price_file)
    except Exception as e:
        st.error(f"Ошибка чтения прайса: {e}")
        st.stop()

    st.write("### Диагностика прайса")
    st.write(price_diag)

    if run_preview:
        preview_df, skipped_df = preview_changes(
            avito_file_obj=avito_file,
            price_map=price_map,
            stock_map=stock_map,
            margin_percent=margin_percent,
            offset_hours=utc_offset_hours,
            active_days=active_days,
            close_days_back=close_days_back,
        )

        st.write("### Что будет обновлено")
        st.write({
            "Всего строк к обработке": len(preview_df),
            "Будут закрыты": int((preview_df["action"] == "Закрыть").sum()) if not preview_df.empty else 0,
            "Будут активны": int((preview_df["action"] == "Оставить активным").sum()) if not preview_df.empty else 0,
            "DateEnd по текущему статусу": int((preview_df["action"] == "DateEnd по статусу").sum()) if not preview_df.empty else 0,
            "Пропущено как неоригинал": len(skipped_df),
        })

        if not preview_df.empty:
            with st.expander("Таблица обновлений"):
                st.dataframe(preview_df, use_container_width=True)

        if not skipped_df.empty:
            with st.expander("Пропущено как неоригинал"):
                st.dataframe(skipped_df, use_container_width=True)

        if stock_file:
            new_ids_df = preview_new_stock_ids(
                stock_file_obj=stock_file,
                avito_file_obj=avito_file,
                price_map=price_map,
            )
            st.write("### Новые ID, которые будут добавлены в stock")
            st.write({"Новых AvitoId": len(new_ids_df)})
            if not new_ids_df.empty:
                with st.expander("Показать новые ID"):
                    st.dataframe(new_ids_df, use_container_width=True)

    if run_prices or run_all:
        try:
            price_output, price_stats = update_prices_and_dateend(
                avito_file_obj=avito_file,
                price_map=price_map,
                stock_map=stock_map,
                margin_percent=margin_percent,
                offset_hours=utc_offset_hours,
                active_days=active_days,
                close_days_back=close_days_back,
            )
            st.success(f"Строк обработано: {price_stats['updated_rows']}")
            st.write({
                "Цены обновлены": price_stats["updated_prices"],
                "DateEnd по количеству": price_stats["dateend_set_from_qty"],
                "DateEnd по текущему статусу": price_stats["dateend_set_from_status"],
                "Не найдено в прайсе": price_stats["unmatched"],
                "Пропущено неоригинала в файле Avito": price_stats["skipped_non_original_avito"],
            })
            with st.expander("Показать примеры обновления"):
                st.write(price_stats["examples"])
            st.download_button(
                "Скачать файл цен + DateEnd",
                data=price_output.getvalue(),
                file_name="avito_prices_dateend_updated_final.xlsx",
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        except Exception as e:
            st.error(f"Ошибка обновления цен/DateEnd: {e}")

    if run_stock or run_all:
        if not stock_file:
            st.error("Для обновления остатков нужен файл остатков.")
        else:
            try:
                stock_output, stock_stats = update_stock_only(
                    stock_file_obj=stock_file,
                    avito_file_obj=avito_file,
                    price_map=price_map,
                    stock_map=stock_map,
                )

                st.success(
                    f"Остатки обновлены. Обновлено строк: {stock_stats['updated_existing']}, "
                    f"добавлено новых ID: {stock_stats['added_new']}"
                )

                st.write("### Отчёт по остаткам")
                st.write({
                    "Stock=0": stock_stats["zero_count"],
                    "Stock=1": stock_stats["one_count"],
                    "Stock=*": stock_stats["many_count"],
                    "Не сопоставлено по артикулу": stock_stats["unmatched_article_count"],
                    "Совпадений объявлений по артикулу": stock_stats["ad_map_stats"]["matched"],
                    "Не нашли артикул в объявлениях": stock_stats["ad_map_stats"]["unmatched"],
                    "Пропущено неоригинала в файле Avito": stock_stats["ad_map_stats"]["skipped_non_original_avito"],
                })

                st.download_button(
                    "Скачать файл остатков",
                    data=stock_output.getvalue(),
                    file_name="avito_stock_only_updated_final.xlsx",
                    mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                )
            except Exception as e:
                st.error(f"Ошибка обновления остатков: {e}")
